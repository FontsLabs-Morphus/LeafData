import os
from airflow.models import BaseOperator
from airflow.hooks.base import BaseHook
from airflow.utils.context import Context

from openpyxl import load_workbook

from custom_operators.metaclass import Source, Destination
from exceptions.custom_exception import MorphusAirflowException


class LoadExcelOperator(BaseOperator, Source):
    """
    Reads an Excel sheet, converts it to CSV through existing Source helpers,
    and pushes the temp CSV path via XCom.
    """

    def __init__(
        self,
        excel_conn_id: str,
        sheet_name: str,
        tx_list: list = None,
        header_less_file: str = "N",
        headers: list = None,
        ext: str = "csv",
        delimiter: str = ",",
        *args,
        **kwargs,
    ):
        super().__init__(*args, **kwargs)
        self.excel_conn_id = excel_conn_id
        self.sheet_name = sheet_name
        self.tx_list = tx_list or []
        self.header_less_file = header_less_file
        self.headers = headers or []
        self.ext = ext
        self.delimiter = delimiter

    def _get_excel_path(self) -> str:
        try:
            conn = BaseHook.get_connection(self.excel_conn_id)

            # For v1, Excel path is stored in the connection's database field.
            excel_path = (conn.schema or "").strip()

            if not excel_path:
                # fallback if your connection creation stores it elsewhere
                excel_path = (conn.host or "").strip()

            if not excel_path:
                raise MorphusAirflowException(
                    message="Excel file path not found in connection",
                    task_id=self.task_id,
                    error_source="LoadExcelOperator._get_excel_path",
                )

            return excel_path

        except MorphusAirflowException:
            raise
        except Exception as e:
            raise MorphusAirflowException(
                message="Failed to resolve Excel connection details",
                task_id=self.task_id,
                error_source="LoadExcelOperator._get_excel_path",
                original_exception=e,
            ) from e

    def _read_excel_rows(self, excel_path: str):
        try:
            if not os.path.exists(excel_path):
                raise MorphusAirflowException(
                    message=f"Excel file not found: {excel_path}",
                    task_id=self.task_id,
                    error_source="LoadExcelOperator._read_excel_rows",
                )

            wb = load_workbook(excel_path, read_only=True, data_only=True)

            if self.sheet_name not in wb.sheetnames:
                raise MorphusAirflowException(
                    message=f"Sheet '{self.sheet_name}' not found in workbook",
                    task_id=self.task_id,
                    error_source="LoadExcelOperator._read_excel_rows",
                )

            ws = wb[self.sheet_name]

            rows = []
            for row in ws.iter_rows(values_only=True):
                cleaned = ["" if cell is None else str(cell) for cell in row]
                rows.append(cleaned)

            if not rows:
                raise MorphusAirflowException(
                    message="Excel sheet is empty",
                    task_id=self.task_id,
                    error_source="LoadExcelOperator._read_excel_rows",
                )

            return rows

        except MorphusAirflowException:
            raise
        except Exception as e:
            raise MorphusAirflowException(
                message="Failed to read Excel workbook",
                task_id=self.task_id,
                error_source="LoadExcelOperator._read_excel_rows",
                original_exception=e,
            ) from e

    def execute(self, context: Context):
        try:
            excel_path = self._get_excel_path()
            self.log.info(f"[EXCEL] Reading file: {excel_path} | sheet: {self.sheet_name}")

            raw_rows = self._read_excel_rows(excel_path)

            source_header = raw_rows[0]
            source_data = raw_rows[1:]

            final_headers = self.headers if self.headers else source_header

            row_iter = self._transform_data_stream(
                raw_iter=iter(source_data),
                header_less_file="yes",
                headers=final_headers,
            )

            self._save_data_stream(
                context=context,
                row_iter=row_iter,
                ext=self.ext,
                delimiter=self.delimiter,
            )

            self.log.info("[EXCEL] Excel data converted and pushed via XCom successfully.")

        except MorphusAirflowException:
            raise
        except Exception as e:
            raise MorphusAirflowException(
                message="Excel source execution failed",
                task_id=self.task_id,
                error_source="LoadExcelOperator.execute",
                original_exception=e,
            ) from e


class PushExcelOperator(BaseOperator, Destination):
    """
    Optional placeholder for future Excel destination support.
    Not needed for EXCEL -> MYSQL v1.
    """

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)

    def execute(self, context: Context):
        raise MorphusAirflowException(
            message="PushExcelOperator is not implemented for version 1",
            task_id=self.task_id,
            error_source="PushExcelOperator.execute",
        )